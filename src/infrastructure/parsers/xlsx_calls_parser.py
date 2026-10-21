from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook

from .schemas import CallRecordSchema, ChargeEntrySchema, ChargesFile

MONEY = Decimal("0.01")


def _month_period(dt) -> tuple[date, date]:
    """Первое и последнее число месяца звонка."""
    first = dt.date().replace(day=1)
    if first.month == 12:
        last = first.replace(year=first.year + 1, month=1, day=1) - __import__("datetime").timedelta(days=1)
    else:
        last = first.replace(month=first.month + 1, day=1) - __import__("datetime").timedelta(days=1)
    return first, last


class XlsxCallsParser:
    """
    Парсер детализации звонков (XLSX, лист «ЛУН»).
    Каждая строка — звонок. Агрегирует в ChargeEntrySchema по нашему номеру и месяцу.
    """

    source_type = "xlsx_calls"

    def __init__(self, sheet_name: str = "ЛУН"):
        self.sheet_name = sheet_name

    def parse(self, content: bytes) -> ChargesFile:
        from io import BytesIO

        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        ws = wb[self.sheet_name] if self.sheet_name in wb.sheetnames else wb.active

        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        # сопоставляем колонки по заголовкам, а не по индексам
        idx = {name: i for i, name in enumerate(header) if name}

        aggregated: dict[tuple[str, date, date], dict] = defaultdict(
            lambda: {"total": Decimal("0"), "duration": 0, "calls": 0, "ban": None}
        )

        for row in rows:
            # пропускаем пустые строки и итоговую =SUM(...)
            if not row or row[idx["Начало вызова"]] is None:
                continue

            record = CallRecordSchema.model_validate(
                {name: row[i] for name, i in idx.items()}
            )

            period_start, period_end = _month_period(record.started_at)
            key = (record.caller_number, period_start, period_end)
            agg = aggregated[key]
            agg["total"] += record.amount
            agg["duration"] += record.duration_sec
            agg["calls"] += 1

        entries: list[ChargeEntrySchema] = []
        for (ctn, period_start, period_end), agg in aggregated.items():
            entries.append(
                ChargeEntrySchema(
                    ctn=ctn,
                    ban=None,
                    ben=None,
                    period_start_date=period_start,
                    period_end_date=period_end,
                    total_amount=agg["total"].quantize(MONEY, rounding=ROUND_HALF_UP),
                    # остальные поля — 0 по дефолту
                )
            )

        return ChargesFile(source_type=self.source_type, entries=entries, summary=None)
